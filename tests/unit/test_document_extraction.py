"""Tests for deterministic local TXT, PDF, and DOCX extraction."""

from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument

from agents.analysis import AnalysisRequest
from backend.schemas import Case, Document
from nlp.extraction import extract_document


def create_pdf(path: Path, pages: list[str]) -> None:
    pdf = fitz.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()


def assert_segment_offsets(result) -> None:  # type: ignore[no-untyped-def]
    for segment in result.segments:
        assert result.text[segment.start_char : segment.end_char] == segment.text


def test_extracts_valid_utf8_txt_and_preserves_text(tmp_path: Path) -> None:
    path = tmp_path / "evidence.txt"
    path.write_text("First line\nSecond line", encoding="utf-8")

    result = extract_document(path, "doc-txt")

    assert result.status == "extracted"
    assert result.text == "First line\nSecond line"
    assert result.segments[0].source.locator == "text:1"
    assert_segment_offsets(result)


def test_extracts_utf8_bom_txt(tmp_path: Path) -> None:
    path = tmp_path / "bom.txt"
    path.write_bytes("\ufeffEvidence text".encode("utf-8"))

    result = extract_document(path, "doc-bom")

    assert result.status == "extracted"
    assert result.text == "Evidence text"


def test_empty_and_invalid_encoded_txt_are_reported(tmp_path: Path) -> None:
    empty_path = tmp_path / "empty.txt"
    empty_path.write_text("", encoding="utf-8")
    invalid_path = tmp_path / "invalid.txt"
    invalid_path.write_bytes(b"\xff\xfe\x00")

    empty = extract_document(empty_path, "doc-empty")
    invalid = extract_document(invalid_path, "doc-invalid")

    assert empty.status == "no_text"
    assert empty.warnings[0].code == "no_extractable_text"
    assert invalid.status == "failed"
    assert invalid.errors[0].code == "invalid_text_encoding"


def test_extracts_single_and_multi_page_pdf_in_page_order(tmp_path: Path) -> None:
    single_path = tmp_path / "single.pdf"
    multi_path = tmp_path / "multi.pdf"
    create_pdf(single_path, ["Single page evidence"])
    create_pdf(multi_path, ["First page", "Second page"])

    single = extract_document(single_path, "doc-single")
    multi = extract_document(multi_path, "doc-multi")

    assert single.status == "extracted"
    assert "Single page evidence" in single.text
    assert multi.status == "extracted"
    assert [segment.source.page for segment in multi.segments] == [1, 2]
    assert multi.segments[0].text.index("First page") >= 0
    assert multi.segments[1].text.index("Second page") >= 0
    assert multi.metadata["page_count"] == 2
    assert_segment_offsets(multi)


def test_blank_pdf_requires_ocr_warning_without_fabricated_text(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    create_pdf(path, [""])

    result = extract_document(path, "doc-blank")

    assert result.status == "no_text"
    assert result.text == ""
    assert {warning.code for warning in result.warnings} == {"no_extractable_text", "ocr_required"}
    assert result.segments[0].source.page == 1


def test_extracts_docx_paragraphs_in_order(tmp_path: Path) -> None:
    path = tmp_path / "paragraphs.docx"
    document = DocxDocument()
    document.add_paragraph("Opening evidence")
    document.add_paragraph("Closing evidence")
    document.save(path)

    result = extract_document(path, "doc-paragraphs")

    assert result.status == "extracted"
    assert [segment.text for segment in result.segments] == ["Opening evidence", "Closing evidence"]
    assert [segment.source.locator for segment in result.segments] == ["paragraph:1", "paragraph:2"]
    assert_segment_offsets(result)


def test_extracts_docx_tables_and_preserves_interleaved_order(tmp_path: Path) -> None:
    path = tmp_path / "interleaved.docx"
    document = DocxDocument()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Date"
    table.cell(0, 1).text = "2024-01-01"
    table.cell(1, 0).text = "Party"
    table.cell(1, 1).text = "Claimant"
    document.add_paragraph("After table")
    document.save(path)

    result = extract_document(path, "doc-table")

    assert result.status == "extracted"
    assert [segment.source.locator for segment in result.segments] == ["paragraph:1", "table:1", "paragraph:2"]
    assert result.segments[1].text == "Date\t2024-01-01\nParty\tClaimant"
    assert result.metadata["table_count"] == 1
    assert_segment_offsets(result)


def test_empty_docx_is_not_reported_as_success(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    DocxDocument().save(path)

    result = extract_document(path, "doc-empty-docx")

    assert result.status == "no_text"
    assert result.warnings[0].code == "no_extractable_text"


@pytest.mark.parametrize(
    ("path_name", "setup", "expected_code"),
    [
        ("unsupported.csv", lambda path: path.write_text("a,b", encoding="utf-8"), "unsupported_file_type"),
        ("corrupt.pdf", lambda path: path.write_bytes(b"not a PDF"), "invalid_pdf"),
        ("corrupt.docx", lambda path: path.write_bytes(b"not a DOCX"), "invalid_docx"),
    ],
)
def test_unsupported_and_corrupted_files_fail_safely(tmp_path: Path, path_name: str, setup, expected_code: str) -> None:  # type: ignore[no-untyped-def]
    path = tmp_path / path_name
    setup(path)

    result = extract_document(path, "doc-invalid")

    assert result.status == "failed"
    assert result.errors[0].code == expected_code
    assert result.text == ""


def test_nonexistent_path_and_directory_are_rejected(tmp_path: Path) -> None:
    missing = extract_document(tmp_path / "missing.txt", "doc-missing")
    directory = extract_document(tmp_path, "doc-directory")

    assert missing.errors[0].code == "file_not_found"
    assert directory.errors[0].code == "directory_not_supported"


def test_file_size_limit_is_enforced(tmp_path: Path) -> None:
    path = tmp_path / "large.txt"
    path.write_text("12345", encoding="utf-8")

    result = extract_document(path, "doc-limit", max_file_size_bytes=4)

    assert result.status == "failed"
    assert result.errors[0].code == "file_too_large"


def test_segment_ordering_and_offsets_are_deterministic(tmp_path: Path) -> None:
    path = tmp_path / "ordered.pdf"
    create_pdf(path, ["One", "Two", "Three"])

    first = extract_document(path, "doc-ordered")
    second = extract_document(path, "doc-ordered")

    assert first.model_dump() == second.model_dump()
    assert [segment.segment_id for segment in first.segments] == [
        "doc-ordered:segment:1",
        "doc-ordered:segment:2",
        "doc-ordered:segment:3",
    ]
    assert_segment_offsets(first)


def test_extracted_text_is_compatible_with_analysis_request(tmp_path: Path) -> None:
    path = tmp_path / "analysis.txt"
    path.write_text("Supplied evidence text", encoding="utf-8")
    extraction = extract_document(path, "doc-analysis")

    request = AnalysisRequest(
        case=Case(id="case-001", title="Example case"),
        document=Document(id=extraction.document_id, case_id="case-001", name=extraction.filename),
        text=extraction.text,
        document_metadata=extraction.metadata,
    )

    assert request.text == "Supplied evidence text"
