"""Actual extraction-to-analysis integration tests using synthetic files."""

from pathlib import Path

import fitz
from docx import Document as DocxDocument
import pytest

from agents.analysis import AnalysisAgent, AnalysisRequest
from backend.schemas import Case, Document
from nlp.extraction import extract_document


def _request(result, case_id: str) -> AnalysisRequest:
    return AnalysisRequest(case=Case(id=case_id, title="Synthetic"), document=Document(id=result.document_id, case_id=case_id, name=result.filename), text=result.text, document_metadata=result.metadata, segments=result.segments)


def test_txt_pdf_and_docx_end_to_end_preserve_provenance(tmp_path: Path) -> None:
    txt = tmp_path / "evidence.txt"; txt.write_text("Nimal Perera paid LKR 5,000 on 12 January 2025.", encoding="utf-8")
    pdf = tmp_path / "evidence.pdf"; document = fitz.open()
    for page_text in ("Nimal Perera paid LKR 5,000 on 12 January 2025.", "Kamal Silva paid LKR 6,000 on 13 January 2025."):
        page = document.new_page(); page.insert_text((72, 72), page_text)
    document.save(pdf); document.close()
    docx = tmp_path / "evidence.docx"; word = DocxDocument(); word.add_paragraph("Opening paragraph")
    table = word.add_table(rows=1, cols=1); table.cell(0, 0).text = "Nimal Perera paid LKR 7,000 on 14 January 2025."; word.save(docx)
    for path, document_id in ((txt, "txt-doc"), (pdf, "pdf-doc"), (docx, "docx-doc")):
        extraction = extract_document(path, document_id)
        response = AnalysisAgent().analyze(_request(extraction, f"case-{document_id}"))
        assert extraction.status == "extracted"
        assert response.status == "completed"
        assert response.model_dump_json()
        for fact in response.facts:
            assert response.document_id == fact.source.document_id
            assert extraction.text[fact.source.start_char:fact.source.end_char] == fact.source.excerpt
    pdf_response = AnalysisAgent().analyze(_request(extract_document(pdf, "pdf-doc"), "case-pdf-doc"))
    assert [fact.source.page for fact in pdf_response.facts] == [1, 2]
    docx_response = AnalysisAgent().analyze(_request(extract_document(docx, "docx-doc"), "case-docx-doc"))
    assert docx_response.facts[0].source.locator == "table:1"
